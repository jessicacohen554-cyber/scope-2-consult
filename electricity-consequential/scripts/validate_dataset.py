#!/usr/bin/env python3
"""Check the built electricity-consequential dataset against the raw export and
against the figures the hub's documentation quotes.

The first checks are the ones that matter for trust: every non-empty source cell
must appear exactly once in responses_long with its value unchanged, and the
multi-select explosion must conserve selections. The rest guard specific figures,
so that editing survey_meta.py or reference/question_labels.csv and rebuilding
cannot silently leave the documentation - or the dashboard - stale.

All frozen counts are on the RAW base of 185 respondents. The analytical base of
181 (PLAN.md gotcha 4) is adjudicated downstream by P20 and enforced by P22; the
exclusion columns are present here but empty by design.

Usage:  python3 electricity-consequential/scripts/validate_dataset.py
        (exit 0 = all pass)
"""
from __future__ import annotations

import csv
import json
import sqlite3
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import survey_meta as meta  # noqa: E402
from build_dataset import DB_NAME, HEADER_PAT, SHEET, SRC, norm_text  # noqa: E402

csv.field_size_limit(10 ** 7)
ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"

# Profile columns. "Substantive" everywhere below means "not one of these".
PROFILE_QNUMS = {3, 4, 5, 6, 9, 10, 11, 12, 13, 14, 15, 16, 17}

# ---------------------------------------------------------------------------
# Frozen answer distributions, read straight off the raw export during planning
# (PLAN.md sections 3.2-3.4) and reproduced here so that a change in option
# matching, whitespace handling or ladder assignment cannot pass unnoticed.
# ---------------------------------------------------------------------------
STANCE_COUNTS = {
    "Q019": {"No": 84, "Yes": 49},
    "Q021": {"Yes": 57, "No": 53},
    "Q024": {"Reported each year": 81,
             "Reported once for the lifetime of the project": 22},
    "Q031": {"Yes": 40, "Unsure, depends on details": 32, "No": 29},
    "Q033": {"Yes": 43, "No": 35},
    "Q043": {"Country": 22, "Grid region": 20, "Nodal": 20,
             "Balancing area": 15, "Zonal": 3},
    "Q045": {"Annual": 31, "Hourly": 30, "Sub-hourly": 14, "Monthly": 6},
}

# question_id -> (n, required, optional, not_required)
MATRIX_COUNTS = {
    "Q026_1": (108, 69, 22, 17),   # Regulatory
    "Q026_2": (106, 52, 31, 23),   # Timing
    "Q026_3": (107, 18, 58, 31),   # Financial analysis
    "Q026_4": (101, 15, 46, 40),   # Barrier
    "Q026_5": (103, 23, 33, 47),   # Common practice
    "Q026_6": (105, 18, 57, 30),   # Positive list
    "Q026_7": (101, 14, 38, 49),   # Performance standard
    "Q026_8": (102, 24, 41, 37),   # Contractual/tenor
    "Q026_9": (99, 10, 36, 53),    # First-of-its-kind
}

SELECTION_COUNTS = {
    "Q028": {"Regulatory test": 80, "Timing test": 72, "Positive list": 58,
             "Contractual/tenor test": 47, "Financial analysis test": 44,
             "Performance standard": 42, "Common practice test": 39,
             "First-of-its-kind test": 35, "Barrier test": 30,
             "None (no tests are feasible)": 10},
    "Q035": {"SCED – locational": 48, "Statistical": 44,
             "SCED – fuel on the margin": 40, "Scenario modeling": 26,
             "Heat-rate/LMP": 23, "Capacity factor based": 20,
             "Difference-based": 18, "None": 7},
    "Q036": {"Capacity factor based": 31, "Difference-based": 28,
             "Scenario modeling": 25, "Heat-rate/LMP": 24,
             "SCED – locational": 12, "SCED – fuel on the margin": 11,
             "Statistical": 10, "None": 9},
    "Q038": {"Recent capacity additions": 48, "Capacity expansion modeling": 30,
             "Policy scenario": 26, "Average emission rate": 17, "None": 7},
    "Q039": {"Average emission rate": 38, "Policy scenario": 29,
             "Capacity expansion modeling": 27, "Recent capacity additions": 11,
             "None": 5},
    "Q047": {"GHG Protocol Guidelines for Quantifying GHG Reductions from "
             "Grid-connected Electricity Projects": 30,
             "UNFCCC CDM Tool07": 19,
             "Default 0.50 build margin weight for all projects": 16,
             "Unsure": 15, "Intervention lifecycle approaches": 13,
             "Resource adequacy approaches": 9, "None are appropriate": 9},
    "Q049": {"Unsure": 21, "Resource adequacy approaches": 17,
             "All are feasible": 16, "Intervention lifecycle approaches": 15,
             "UNFCCC CDM Tool07": 10,
             "Default 0.50 build margin weight for all projects": 8,
             "GHG Protocol Guidelines for Quantifying GHG Reductions from "
             "Grid-connected Electricity Projects": 8},
}

# The survey's attrition spine (PLAN.md 3.1). Response counts fall from 165 to
# 36; every percentage anywhere downstream has to divide by the right one.
ATTRITION = {"Q018": 165, "Q019": 133, "Q021": 110, "Q024": 103, "Q028": 98,
             "Q035": 69, "Q043": 80, "Q045": 81, "Q047": 67, "Q052": 36}

# ---------------------------------------------------------------------------
# The P02 editorial layer, frozen. A rebuild that shifts any of these has
# either edited the codebook deliberately (update the figures here in the same
# commit) or broken the label join.
# ---------------------------------------------------------------------------
TOPIC_COUNTS = {"profile": 13, "general": 1, "formula": 7, "additionality": 17,
                "emission_rates": 12, "weighting": 6}
CATEGORY_COUNTS = {"resp_profile": 13, "overall_assessment": 1,
                   "quantification_design": 11, "additionality_design": 11,
                   "feasibility_and_data": 4, "method_choice": 11,
                   "regional_variation": 2, "claims_and_rigor": 2,
                   "evidence": 1}
ASKS_FOR_COUNTS = {"respondent_attribute": 13, "stance": 5, "matrix_rating": 9,
                   "feasibility_pick": 2, "method_pick": 5,
                   "design_preference": 2, "rationale": 10, "elaboration": 4,
                   "evidence": 1, "open_feedback": 5}
DOC_SECTION_COUNTS = {"n/a": 13, "5": 1, "6.1": 7, "7.1": 17, "8.1": 3,
                      "8.2": 3, "8": 6, "9.1": 6}
# Critical answer per single_select stance question, proposal-relative (see
# survey_meta.POLARITY for the bases). None = judged to have no drafted
# position; 21 is deliberately "Yes" - the draft excludes secondary effects.
STANCE_POLARITY = ((19, "No"), (21, "Yes"),
                   (24, "Reported once for the lifetime of the project"),
                   (31, None), (33, None))

# The junk/test respondents P10 detects and P20 adjudicates. Frozen here only as
# a cell count, because P01 does not exclude anybody.
JUNK_CANDIDATES = (11, 12, 14)

VIEWS = ("v_stance_answers", "v_stance_summary", "v_selections", "v_free_text",
         "v_option_counts", "v_answer_types", "v_question_tree",
         "v_stance_by_redaction", "v_redaction_profile")


def main() -> int:
    wb = openpyxl.load_workbook(SRC, read_only=True, data_only=True)
    ws = wb[SHEET]
    it = ws.iter_rows(values_only=True)
    header = [norm_text(h) for h in next(it)]
    raw_rows = [list(r) for r in it]
    raw_rows = [r for r in raw_rows if any(norm_text(v) for v in r)]
    rows = [[norm_text(v) for v in r] for r in raw_rows]
    wb.close()

    qidx, qid_of, qnum_of = [], {}, {}
    for i, h in enumerate(header):
        m = HEADER_PAT.match(h)
        if not m:
            continue
        qnum = int(m.group(1))
        sub = int(m.group(2)) if m.group(2) else None
        qidx.append(i)
        qid_of[i] = f"Q{qnum:03d}" if sub is None else f"Q{qnum:03d}_{sub}"
        qnum_of[i] = qnum

    raw_cells = sum(1 for r in rows for i in qidx if r[i])

    with (DATA / "responses_long.csv").open(encoding="utf-8") as f:
        long_rows = list(csv.DictReader(f))
    with (DATA / "response_selections.csv").open(encoding="utf-8") as f:
        sel_rows = list(csv.DictReader(f))
    with (DATA / "questions.csv").open(encoding="utf-8") as f:
        q_rows = list(csv.DictReader(f))
    manifest = json.loads((DATA / "manifest.json").read_text(encoding="utf-8"))

    # Free-text-ness is a property of the built codebook, not of the raw grid,
    # so the raw-side figures are derived from the codebook and then counted on
    # the raw cells - which is what makes them a check rather than a tautology.
    free_text_qids = {q["question_id"] for q in q_rows
                      if q["question_type"] == "free_text"}
    raw_substantive_free_text = sum(
        1 for r in rows for i in qidx
        if r[i] and qid_of[i] in free_text_qids and qnum_of[i] not in PROFILE_QNUMS)
    # The survey tool's 4,000-character cap, measured on the SOURCE cell: 18
    # substantive answers land in the top 10 characters and the next longest is
    # 3,986, so the band is a cliff rather than a cut-off we chose.
    raw_capped = sum(
        1 for r in raw_rows for i in qidx
        if r[i] and qid_of[i] in free_text_qids and qnum_of[i] not in PROFILE_QNUMS
        and len(str(r[i])) >= meta.TEXT_LIMIT - meta.TRUNCATION_BAND)

    lookup = {(int(r["respondent_id"]), r["question_id"]): r["answer_text"]
              for r in long_rows}
    mismatches = sum(
        1 for r in rows for i in qidx
        if r[i] and lookup.get((int(r[0]), qid_of[i])) != r[i])
    n_selected = sum(int(r["n_selected"]) for r in long_rows if r["n_selected"])
    junk_cells = sum(1 for r in long_rows
                     if int(r["respondent_id"]) in JUNK_CANDIDATES)

    con = sqlite3.connect(DATA / DB_NAME)
    scalar = lambda sql, *a: con.execute(sql, a).fetchone()[0]  # noqa: E731

    def option_counts(qid):
        return {t: n for t, n in con.execute(
            "SELECT answer_text, COUNT(*) FROM responses WHERE question_id = ? "
            "GROUP BY answer_text", (qid,))}

    def selection_counts(qid):
        return {t: n for t, n in con.execute(
            "SELECT option_text, COUNT(*) FROM response_selections "
            "WHERE question_id = ? GROUP BY option_text", (qid,))}

    def label_dist(col):
        return {v: n for v, n in con.execute(
            f"SELECT {col}, COUNT(*) FROM questions GROUP BY {col}")}

    matrix_got = {
        qid: (scalar("SELECT COUNT(*) FROM responses WHERE question_id = ?", qid),)
             + tuple(option_counts(qid).get(lvl, 0)
                     for lvl in ("Required", "Optional", "Not required"))
        for qid in MATRIX_COUNTS
    }

    checks = [
        # ---- integrity: the dataset is the export, reshaped and nothing else --
        ("round-trip cell count", raw_cells, len(long_rows)),
        ("round-trip value mismatches", 0, mismatches),
        ("selection explosion conserved", n_selected, len(sel_rows)),
        ("sqlite responses match csv", len(long_rows),
         scalar("SELECT COUNT(*) FROM responses")),
        ("sqlite selections match csv", len(sel_rows),
         scalar("SELECT COUNT(*) FROM response_selections")),
        ("sqlite questions match csv", len(q_rows),
         scalar("SELECT COUNT(*) FROM questions")),
        ("manifest agrees with responses_long", manifest["rows"]["responses_long"],
         len(long_rows)),
        # ---- structure -----------------------------------------------------
        ("every question has an anchor or is one", 0, scalar(
            "SELECT COUNT(*) FROM questions WHERE anchor_question IS NULL "
            "AND role NOT IN ('profile','other_specify')")),
        ("every follow-up has a parent", 0, scalar(
            "SELECT COUNT(*) FROM questions WHERE parent_question IS NULL "
            "AND role IN ('comment','other_specify')")),
        ("no parent points forward", 0, scalar(
            "SELECT COUNT(*) FROM questions "
            "WHERE parent_question >= question_number")),
        ("every parent exists", 0, scalar(
            "SELECT COUNT(*) FROM questions c WHERE c.parent_question IS NOT NULL "
            "AND NOT EXISTS (SELECT 1 FROM questions p "
            "WHERE p.question_number = c.parent_question)")),
        ("matrix carried as 9 sub-numbered rows", 9, scalar(
            "SELECT COUNT(*) FROM questions WHERE question_number = 26 "
            "AND sub_number BETWEEN 1 AND 9")),
        ("matrix ranks are the stringency ladder", 0, scalar(
            "SELECT COUNT(*) FROM responses WHERE question_type = 'matrix_rating' "
            "AND answer_numeric NOT BETWEEN 1 AND 3")),
        ("ordered selects ranked coarse-to-fine", 0, scalar(
            "SELECT COUNT(*) FROM responses WHERE question_type = 'ordinal_select' "
            "AND answer_numeric IS NULL")),
        ("special options carry no rank", 0, scalar(
            "SELECT COUNT(*) FROM question_options "
            "WHERE is_special = 1 AND option_rank IS NOT NULL")),
        ("specials flagged on the picklists", 12, scalar(
            "SELECT COUNT(*) FROM question_options WHERE is_special = 1")),
        # ---- hand-curated labels -------------------------------------------
        ("all questions labelled", 56, scalar(
            "SELECT COUNT(*) FROM questions WHERE shorthand IS NOT NULL "
            "AND label IS NOT NULL AND topic IS NOT NULL "
            "AND doc_section IS NOT NULL AND category IS NOT NULL "
            "AND subcategory IS NOT NULL AND asks_for IS NOT NULL")),
        ("shorthands unique", 56, scalar(
            "SELECT COUNT(DISTINCT shorthand) FROM questions")),
        ("shorthands snake_case within 40 chars", 0, scalar(
            "SELECT COUNT(*) FROM questions WHERE LENGTH(shorthand) > 40 "
            "OR shorthand GLOB '*[^a-z0-9_]*' OR shorthand GLOB '[^a-z]*'")),
        ("labels reach every answer", len(long_rows), scalar(
            "SELECT COUNT(*) FROM responses WHERE shorthand IS NOT NULL "
            "AND topic IS NOT NULL AND asks_for IS NOT NULL")),
        ("answer-type counts conserve answers", len(long_rows), scalar(
            "SELECT SUM(n_answers) FROM v_answer_types")),
        # ---- the P02 editorial layer ---------------------------------------
        ("labels final, no provisional placeholders", 0, scalar(
            "SELECT COUNT(*) FROM questions WHERE ? IN "
            "(topic, doc_section, category, subcategory, asks_for)",
            meta.PROVISIONAL_LABEL_VALUE)),
        ("provisional labels disallowed in meta", False,
         meta.ALLOW_PROVISIONAL_LABELS),
        ("shorthands carry no question numbers", 0, scalar(
            "SELECT COUNT(*) FROM questions WHERE shorthand GLOB '*[0-9]*'")),
        ("topic distribution", TOPIC_COUNTS, label_dist("topic")),
        ("category distribution", CATEGORY_COUNTS, label_dist("category")),
        ("asks_for distribution", ASKS_FOR_COUNTS, label_dist("asks_for")),
        ("doc_section distribution", DOC_SECTION_COUNTS,
         label_dist("doc_section")),
        ("subcategories inside the closed pairing", (23, 0),
         (scalar("SELECT COUNT(DISTINCT subcategory) FROM questions"),
          sum(1 for cat, sub in con.execute(
              "SELECT DISTINCT category, subcategory FROM questions")
              if sub not in meta.CATEGORY_SUBCATS.get(cat, set())))),
        ("polarity covers each stance question", STANCE_POLARITY, tuple(
            (q, (meta.POLARITY[q] or {}).get("critical")
                if q in meta.POLARITY else "MISSING")
            for q, in con.execute(
                "SELECT question_number FROM questions "
                "WHERE question_type = 'single_select' AND asks_for = 'stance' "
                "ORDER BY question_number"))),
        ("polarity names real answer options", 0, sum(
            1 for q, entry in meta.POLARITY.items() if entry
            for side in ("critical", "supportive")
            if entry[side] not in option_counts(f"Q{q:03d}"))),
        ("hazard notes present where required", 0, scalar(
            "SELECT COUNT(*) FROM questions WHERE question_id IN (%s) "
            "AND (label_notes IS NULL OR label_notes = '')"
            % ",".join("?" * len(meta.REQUIRED_NOTES)),
            *sorted(meta.REQUIRED_NOTES))),
        ("ladder anchors cover the ordinal ladders",
         sorted(meta.ORDINAL_LADDERS), sorted(meta.ORDINAL_ANCHORS)),
        # ---- frozen figures for this dataset --------------------------------
        ("respondents", 185, scalar("SELECT COUNT(*) FROM respondents")),
        ("respondent ids run 3..201 with gaps", "3..201/185",
         "%d..%d/%d" % (scalar("SELECT MIN(respondent_id) FROM respondents"),
                        scalar("SELECT MAX(respondent_id) FROM respondents"),
                        scalar("SELECT COUNT(*) FROM respondents"))),
        ("question columns", 56, scalar("SELECT COUNT(*) FROM questions")),
        ("profile questions", 13, scalar(
            "SELECT COUNT(*) FROM questions WHERE question_number <= 17")),
        ("substantive questions (34 numbered + 9 matrix)", 43, scalar(
            "SELECT COUNT(*) FROM questions WHERE question_number >= 18")),
        ("non-empty answer cells", 5074, len(long_rows)),
        ("substantive free-text answers", 1377, raw_substantive_free_text),
        ("substantive free-text answers in db", 1377, scalar(
            "SELECT COUNT(*) FROM v_free_text WHERE question_number >= 18")),
        ("exploded selections", 1277, len(sel_rows)),
        ("write-in selections", 0,
         sum(1 for r in sel_rows if r["is_canonical"] == "0")),
        ("redacted respondents flagged", 77, scalar(
            "SELECT COUNT(*) FROM respondents WHERE is_redacted = 1")),
        ("named respondents", 108, scalar(
            "SELECT COUNT(*) FROM respondents WHERE is_redacted = 0")),
        ("redacted labelled, not nulled", 77, scalar(
            "SELECT COUNT(*) FROM respondents "
            "WHERE name = 'Redacted' AND organization = 'Redacted'")),
        ("no null name or organization", 0, scalar(
            "SELECT COUNT(*) FROM respondents "
            "WHERE name IS NULL OR organization IS NULL")),
        ("exclusion columns present and empty", 0, scalar(
            "SELECT COUNT(*) FROM respondents "
            "WHERE is_excluded <> 0 OR exclusion_reason IS NOT NULL")),
        ("junk-candidate cells (ids 11, 12, 14)", 125, junk_cells),
        ("cap-length substantive answers", 18, raw_capped),
        ("cap-length flag matches the raw cliff", raw_capped,
         sum(1 for r in long_rows if r["likely_truncated"] == "1"
             and int(r["question_number"]) >= 18)),
        ("respondents answering nothing substantive", 2, scalar(
            "SELECT COUNT(*) FROM respondents WHERE n_substantive_answered = 0")),
        ("views all return rows", len(VIEWS), sum(
            1 for v in VIEWS if scalar(f"SELECT COUNT(*) FROM {v}") > 0)),
    ]

    # ---- the distributions every downstream page is built on ---------------
    for qid, want in STANCE_COUNTS.items():
        checks.append((f"{qid} answer split", want, option_counts(qid)))
    for qid, want in MATRIX_COUNTS.items():
        checks.append((f"{qid} n/R/O/N", want, matrix_got[qid]))
    for qid, want in SELECTION_COUNTS.items():
        checks.append((f"{qid} selection tallies", want, selection_counts(qid)))
    for qid, want in ATTRITION.items():
        checks.append((f"{qid} respondents answering", want, scalar(
            "SELECT n_answered FROM questions WHERE question_id = ?", qid)))

    con.close()

    failed = 0
    for name, expected, got in checks:
        good = expected == got
        failed += not good
        line = f"  {'PASS' if good else 'FAIL'}  {name:<46}"
        if isinstance(expected, dict):
            print(f"{line} n={sum(expected.values())}"
                  + ("" if good else f"\n        expected={expected}\n        got={got}"))
        else:
            print(f"{line} expected={expected!s:<12} got={got}")
    print(f"\n{len(checks) - failed}/{len(checks)} checks passed")
    return 1 if failed else 0


if __name__ == "__main__":
    sys.exit(main())
